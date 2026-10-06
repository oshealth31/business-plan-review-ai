"""Tests for the pieces specific to FundingLens Check that need no web stack:
templates, applicant report, billing/webhook helpers, documents, AI wrapper, printable report.

    python -m unittest tests.test_check_logic -v
"""
import io
import json
import os
import sys
import unittest
from datetime import datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from fundinglens_core import analysis as A  # noqa: E402
from fundinglens_core import documents as D  # noqa: E402
from fundinglens_core import readiness as R  # noqa: E402
from fundinglens_core import templates as T  # noqa: E402
from fundinglens_core.ai import ai_commentary  # noqa: E402
from backend import billing as B  # noqa: E402
from backend.report_html import render_report_html  # noqa: E402

DEMO = open(os.path.join(ROOT, 'demo_data', 'business_plan_demo.txt'), encoding='utf-8').read()


def run(text, template_id='sme-growth-general', funding=500_000, contribution=100_000, overrides=None, ai=None):
    t = T.get_template(template_id)
    det = A.analyse_document(text, funding, contribution, T.criteria_for(t), overrides)
    return t, A.build_result(det, ai, {'mode': 'rules_only', 'message': 'rules only'})


class TemplateTests(unittest.TestCase):
    def test_templates_are_internally_valid(self):
        self.assertEqual(T.validate_templates(), [])

    def test_every_template_is_honestly_labelled(self):
        for t in T.TEMPLATES:
            self.assertIn(t['basis'], ('generic', 'published'))
            if t['basis'] == 'generic':
                self.assertIn('not the criteria of any specific funder', t['source_note'])
            else:
                self.assertIn('not produced, reviewed or endorsed', t['source_note'])
                self.assertTrue(t['sources'] and t['retrieved'] and t['funder'])
                self.assertEqual(t['review_status'], 'needs_human_review')
                self.assertFalse([c for c in t['criteria'] if c['kind'] == 'numeric'])

    def test_public_view_hides_weights(self):
        for v in T.list_public():
            blob = json.dumps(v)
            self.assertNotIn('weight', blob)
            self.assertTrue(v['criteria'])

    def test_requirement_text_is_plain_language(self):
        t = T.get_template('sme-growth-general')
        crits = {c['key']: T.requirement_text(c) for c in t['criteria']}
        self.assertEqual(crits['operating_margin'], 'Operating margin at least 15%')
        self.assertEqual(crits['funding_to_contribution'], 'Funding compared with your contribution at most 5x')
        self.assertTrue(crits['market_demand'].startswith('Your plan should cover'))

    def test_criteria_for_is_a_copy(self):
        t = T.get_template('sme-growth-general')
        c = T.criteria_for(t)
        c[0]['threshold'] = 999
        self.assertNotEqual(T.get_template('sme-growth-general')['criteria'][0]['threshold'], 999)

    def test_every_template_runs_through_the_engine(self):
        for t in T.TEMPLATES:
            det = A.analyse_document(DEMO, 500_000, 100_000, T.criteria_for(t))
            self.assertEqual(len(det['rubric']['criteria']), len(t['criteria']))

    def test_unknown_template(self):
        self.assertIsNone(T.get_template('nope'))


class ReadinessTests(unittest.TestCase):
    def test_demo_plan_has_no_required_gap_on_general_template(self):
        t, res = run(DEMO)
        rep = R.applicant_report(res, t)
        self.assertEqual(rep['summary']['required_outstanding'], 0)
        self.assertIn(rep['summary']['band'], ('well_covered', 'mostly_covered'))

    def test_weak_plan_needs_work_and_gets_advice(self):
        t, res = run('We plan to open a shop.', funding=900_000, contribution=10_000)
        rep = R.applicant_report(res, t)
        self.assertEqual(rep['summary']['band'], 'needs_work')
        gaps = [i for i in rep['items'] if i['state'] in ('not_met', 'missing')]
        self.assertTrue(gaps)
        for g in gaps:
            self.assertTrue(g['how_to_fix'] and g['reviewer_asks'], g['key'])
        # required gaps are listed first
        first = rep['items'][0]
        self.assertTrue(first['required'] and first['state'] in ('not_met', 'missing'))

    def test_met_items_have_no_fix_text(self):
        t, res = run(DEMO)
        met = [i for i in R.applicant_report(res, t)['items'] if i['state'] == 'met']
        self.assertTrue(met)
        self.assertTrue(all(i['how_to_fix'] == '' for i in met))

    def test_mentioned_is_not_presented_as_proven(self):
        t, res = run(DEMO)
        m = [i for i in R.applicant_report(res, t)['items'] if i['state'] == 'mentioned']
        self.assertTrue(m)
        self.assertTrue(all('not proof' in i['detail'] for i in m))
        self.assertTrue(all(i['how_to_fix'] for i in m))  # still tell them how to back it up

    def test_negated_is_a_gap(self):
        t, res = run('No market research has been done. Annual revenue: BWP 100,000. Annual expenses: BWP 50,000.')
        item = [i for i in R.applicant_report(res, t)['items'] if i['key'] == 'market_demand'][0]
        self.assertEqual(item['state'], 'not_met')
        self.assertIn('missing or not yet done', item['detail'])

    def test_missing_figures_make_the_report_incomplete_not_failed(self):
        t, res = run('Market demand is strong. Competitors are few. Cash flow is attached.', funding=0, contribution=0)
        rep = R.applicant_report(res, t)
        states = {i['key']: i['state'] for i in rep['items']}
        self.assertEqual(states['operating_margin'], 'cannot_assess')
        self.assertIn(rep['summary']['band'], ('incomplete', 'needs_work'))

    def test_no_weights_or_percentages_leak(self):
        t, res = run(DEMO)
        blob = json.dumps(R.applicant_report(res, t))
        self.assertNotIn('"weight"', blob)
        self.assertNotIn('coverage_pct', blob)

    def test_quick_scan_reveals_counts_only(self):
        t, res = run('We plan to open a shop.')
        q = R.quick_scan(res)
        blob = json.dumps(q)
        for c in t['criteria']:
            self.assertNotIn(c['label'], blob)
        self.assertNotIn('how_to_fix', blob)
        self.assertEqual(q['counts']['total'], len(t['criteria']))
        self.assertIn('Unlock the full report', q['message'])

    def test_quick_scan_lists_figures_to_confirm(self):
        _, res = run(DEMO)
        self.assertEqual(sorted(R.quick_scan(res)['figures_to_confirm']), ['expenses', 'revenue'])

    def test_confirmed_figures_stop_asking(self):
        _, res = run(DEMO, overrides={'revenue': 1_200_000, 'expenses': 900_000})
        self.assertEqual(R.quick_scan(res)['figures_to_confirm'], [])

    def test_flags_exclude_items_already_in_checklist(self):
        t, res = run('We plan to open a shop.')
        rep = R.applicant_report(res, t)
        titles = [f['title'] for f in rep['flags']]
        self.assertFalse(any('below standard' in x for x in titles))
        self.assertTrue(any('Revenue figure not identified' in x for x in titles))

    def test_ai_overview_only_when_ai_ran(self):
        t, res = run(DEMO)
        self.assertEqual(R.applicant_report(res, t)['overview'], '')
        res['ai'] = {'mode': 'llm', 'message': 'ok'}
        res['executive_summary'] = 'A processing business.'
        self.assertEqual(R.applicant_report(res, t)['overview'], 'A processing business.')

    def test_disclaimer_present_and_never_promises(self):
        t, res = run(DEMO)
        rep = R.applicant_report(res, t)
        self.assertIn('not a prediction', rep['disclaimer'])
        for banned in ('likely to be approved', 'will be approved', 'guarantee'):
            self.assertNotIn(banned, json.dumps(rep).lower().replace('does not guarantee', ''))

    def test_band_logic(self):
        self.assertEqual(R.readiness_band([('met', True), ('met', False)])['band'], 'well_covered')
        self.assertEqual(R.readiness_band([('met', True)] * 4 + [('missing', False)])['band'], 'mostly_covered')
        self.assertEqual(R.readiness_band([('met', False), ('missing', False), ('missing', False)])['band'], 'needs_work')
        self.assertEqual(R.readiness_band([('missing', True), ('met', False)])['band'], 'needs_work')
        self.assertEqual(R.readiness_band([('cannot_assess', True), ('met', False)])['band'], 'incomplete')
        self.assertEqual(R.readiness_band([('judgement', False)])['band'], 'incomplete')


class BillingTests(unittest.TestCase):
    P = B.Pricing(full=2, recheck=1, window_days=30)
    NOW = datetime(2026, 10, 5, 12, 0, 0)

    def test_first_report_full_price(self):
        self.assertEqual(B.price_for_full(self.P, None, self.NOW), 2)

    def test_recheck_inside_window_is_cheaper(self):
        self.assertEqual(B.price_for_full(self.P, self.NOW - timedelta(days=29), self.NOW), 1)
        self.assertEqual(B.price_for_full(self.P, self.NOW - timedelta(days=30), self.NOW), 1)

    def test_after_window_full_price_again(self):
        self.assertEqual(B.price_for_full(self.P, self.NOW - timedelta(days=31), self.NOW), 2)

    def test_balance(self):
        self.assertEqual(B.balance([5, -2, -1, 3]), 5)
        self.assertEqual(B.balance([]), 0)

    def test_webhook_signature_roundtrip(self):
        body = b'{"event_id":"evt_1"}'
        sig = B.sign_webhook('s3cret', 1_000, body)
        self.assertEqual(B.verify_webhook('s3cret', 1000, sig, body, now=1_100), (True, 'ok'))

    def test_webhook_rejections(self):
        body = b'{"a":1}'
        sig = B.sign_webhook('s3cret', 1000, body)
        self.assertFalse(B.verify_webhook('wrong', 1000, sig, body, now=1000)[0])
        self.assertFalse(B.verify_webhook('s3cret', 1000, sig, body + b' ', now=1000)[0])           # body tampered
        self.assertEqual(B.verify_webhook('s3cret', 1000, sig, body, now=2000)[1], 'timestamp outside tolerance')  # replay
        self.assertEqual(B.verify_webhook('s3cret', 'x', sig, body)[1], 'bad timestamp')
        self.assertEqual(B.verify_webhook('', 1000, sig, body)[1], 'webhook is not configured')
        self.assertFalse(B.verify_webhook('s3cret', 1000, None, body, now=1000)[0])

    def test_webhook_signature_is_case_insensitive_hex(self):
        sig = B.sign_webhook('k', 5, b'x')
        self.assertTrue(B.verify_webhook('k', 5, sig.upper(), b'x', now=5)[0])

    def test_event_parsing(self):
        good = json.dumps({'event_id': 'evt_123', 'type': 'credits.purchased', 'email': ' Ana@Example.com ', 'credits': 10,
                           'amount': '120.00', 'currency': 'BWP'}).encode()
        e = B.parse_webhook_event(good)
        self.assertEqual((e['email'], e['credits'], e['event_id']), ('ana@example.com', 10, 'evt_123'))
        for bad in (b'nope', b'[]',
                    json.dumps({'event_id': 'x', 'type': 'credits.purchased', 'email': 'a@b.co', 'credits': 1}).encode(),
                    json.dumps({'event_id': 'evt_1', 'type': 'refund', 'email': 'a@b.co', 'credits': 1}).encode(),
                    json.dumps({'event_id': 'evt_1', 'type': 'credits.purchased', 'email': 'not-an-email', 'credits': 1}).encode(),
                    json.dumps({'event_id': 'evt_1', 'type': 'credits.purchased', 'email': 'a@b.co', 'credits': 0}).encode(),
                    json.dumps({'event_id': 'evt_1', 'type': 'credits.purchased', 'email': 'a@b.co', 'credits': True}).encode(),
                    json.dumps({'event_id': 'evt_1', 'type': 'credits.purchased', 'email': 'a@b.co', 'credits': 1.5}).encode()):
            with self.assertRaises(ValueError, msg=bad):
                B.parse_webhook_event(bad)

    def test_packs(self):
        raw = json.dumps([{'id': 'starter', 'name': 'Starter', 'credits': 6, 'price': '150', 'currency': 'BWP', 'url': 'https://pay.example/x'}])
        self.assertEqual(B.parse_packs(raw)[0]['credits'], 6)
        self.assertEqual(B.parse_packs(''), [])
        self.assertEqual(B.parse_packs(None), [])
        for bad in ('{"a":1}', '[{"id":"A!","credits":1}]', '[{"id":"ok","credits":0}]',
                    '[{"id":"ok","credits":1,"url":"http://insecure"}]', '["x"]'):
            with self.assertRaises(ValueError, msg=bad):
                B.parse_packs(bad)

    def test_invite_codes(self):
        codes = {B.make_invite_code() for _ in range(200)}
        self.assertEqual(len(codes), 200)
        for c in codes:
            self.assertRegex(c, r'^FL-[A-Z2-9]{4}-[A-Z2-9]{4}$')
            self.assertFalse(set('01OIL') & set(c.replace('FL-', '')))
        self.assertEqual(B.normalise_invite_code(' fl-ab cd-2345 '), 'FL-ABCD-2345')

    def test_email_helpers(self):
        self.assertEqual(B.normalise_email('  A@B.Co '), 'a@b.co')
        self.assertTrue(B.valid_email('a@b.co'))
        for bad in ('', 'a@b', 'a b@c.de', '@b.co', 'a@' + 'x' * 260 + '.com'):
            self.assertFalse(B.valid_email(bad), bad)

    def test_password_policy(self):
        self.assertIsNotNone(B.password_problem('short1', 'a@b.co'))
        self.assertIsNotNone(B.password_problem('aaaaaaaaaaaaaaaa', 'a@b.co'))
        self.assertIsNotNone(B.password_problem('xxxmelissaxxx99', 'melissa@b.co'))
        self.assertIsNone(B.password_problem('correct horse battery', 'melissa@b.co'))

    def test_retention_deadline(self):
        self.assertEqual(B.retention_deadline(self.NOW, 30), self.NOW + timedelta(days=30))


class DocumentTests(unittest.TestCase):
    def test_text_file(self):
        self.assertIn('revenue', D.extract_text(b'Annual revenue: BWP 100,000 and more text to pass the minimum length check.', '.txt'))

    def test_rejections(self):
        for data, suffix, status in ((b'x', '.exe', 415), (b'', '.txt', 400), (b'not pdf', '.pdf', 415), (b'not zip', '.docx', 415),
                                     (b'short', '.txt', 422)):
            with self.assertRaises(D.DocumentError) as cm:
                D.extract_text(data, suffix)
            self.assertEqual(cm.exception.status, status, (suffix, cm.exception.message))

    def test_docx_roundtrip_including_tables(self):
        from docx import Document
        d = Document()
        d.add_paragraph('Annual revenue: BWP 1,200,000. The market for our product is growing steadily in the region.')
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = 'Expenses', '900,000'
        buf = io.BytesIO()
        d.save(buf)
        text = D.extract_text(buf.getvalue(), '.docx')
        self.assertIn('1,200,000', text)
        self.assertIn('Expenses | 900,000', text)

    def test_corrupt_docx(self):
        with self.assertRaises(D.DocumentError):
            D.extract_text(b'PK\x03\x04 this is not really a zip file', '.docx')

    def test_empty_text_pdf_is_reported_as_scan(self):
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        buf = io.BytesIO()
        w.write(buf)
        with self.assertRaises(D.DocumentError) as cm:
            D.extract_text(buf.getvalue(), '.pdf')
        self.assertEqual(cm.exception.status, 422)
        self.assertIn('Scanned', cm.exception.message)


class _FakeClient:
    def __init__(self, payload=None, boom=False):
        self.payload, self.boom, self.calls = payload, boom, []
        self.responses = self

    def create(self, **kw):
        self.calls.append(kw)
        if self.boom:
            raise TimeoutError('slow')
        return type('R', (), {'output_text': json.dumps(self.payload)})()


class AiWrapperTests(unittest.TestCase):
    def setUp(self):
        self.t = T.get_template('sme-growth-general')
        self.crit = T.criteria_for(self.t)
        self.det = A.analyse_document(DEMO, 500_000, 100_000, self.crit)

    def call(self, client=None, key='k', model='m'):
        return ai_commentary(api_key=key, model=model, timeout=5, text=DEMO, applicant='Applicant', det=self.det,
                             criteria=self.crit, client=client)

    def test_not_configured_is_rules_only(self):
        out, meta, raw = self.call(key=None)
        self.assertIsNone(out)
        self.assertEqual(meta['mode'], 'rules_only')
        self.assertIsNone(raw)

    def test_success_sanitises_and_verifies_quotes(self):
        payload = {'executive_summary': 'Processing business', 'business_model': 'B2B', 'strengths': ['x'], 'assumptions': [],
                   'risks': [{'severity': 'weird', 'title': 'Concentration', 'detail': 'd', 'quote': 'recurring demand, supplier quotations'}],
                   'evidence': [{'status': 'supported', 'claim': 'Invented', 'detail': '', 'quote': 'This sentence is not in the plan at all'}],
                   'criterion_notes': [{'key': 'business_viability', 'note': 'plausible', 'quote': ''}], 'questions': ['q?']}
        client = _FakeClient(payload)
        out, meta, raw = self.call(client)
        self.assertEqual(meta['mode'], 'llm')
        self.assertEqual(out['risks'][0]['severity'], 'medium')
        self.assertTrue(out['risks'][0]['quote_verified'])
        self.assertEqual(out['evidence'][0]['status'], 'needs_verification')
        self.assertIsNotNone(raw)
        sent = client.calls[0]['input']
        self.assertIn('<business_plan>', sent)
        self.assertNotIn('Acme', sent)

    def test_failure_never_raises_and_is_flagged(self):
        out, meta, raw = self.call(_FakeClient(boom=True))
        self.assertIsNone(out)
        self.assertEqual(meta['mode'], 'rules_only')
        self.assertTrue(meta.get('failed'))

    def test_garbage_output_degrades(self):
        out, meta, _ = self.call(_FakeClient(payload='not an object'))
        self.assertIsNone(out)
        self.assertTrue(meta.get('failed'))


class ReportHtmlTests(unittest.TestCase):
    def test_everything_is_escaped(self):
        t, res = run('We plan to open a shop.')
        res['strengths'] = ['<script>alert(1)</script>']
        rep = R.applicant_report(res, t)
        rep['figures'][0]['source_line'] = '<img src=x onerror=alert(1)>'
        html = render_report_html('<b>My plan</b>', rep, generated='2026-10-05', version='0.1.0')
        self.assertNotIn('<script>alert(1)</script>', html)
        self.assertNotIn('<img src=x', html)
        self.assertIn('&lt;b&gt;My plan&lt;/b&gt;', html)
        self.assertIn('Readiness report', html)
        self.assertIn('not a prediction', html)

    def test_contains_fix_advice(self):
        t, res = run('We plan to open a shop.')
        html = render_report_html('Shop', R.applicant_report(res, t), generated='x', version='0.1.0')
        self.assertIn('How to fix', html)
        self.assertIn('A reviewer may ask', html)


if __name__ == '__main__':
    unittest.main()


class PublishedPackTests(unittest.TestCase):
    def test_pack_report_has_attachments_sources_and_renders(self):
        from backend.report_html import render_report_html
        t, res = run(DEMO, 'ceda-agri')
        rep = R.applicant_report(res, t)
        self.assertTrue(rep['documents_to_prepare'] and rep['confirm_yourself'] and rep['template']['sources'])
        self.assertNotIn('weight', json.dumps(T.public_view(t)))
        html = render_report_html('x', rep, generated='now', version='0')
        self.assertIn('Documents to prepare', html)
        self.assertIn('Sources (read 2026-10-05)', html)
        self.assertNotIn('<script', html)
