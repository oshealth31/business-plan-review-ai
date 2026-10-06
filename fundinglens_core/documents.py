"""Document validation and text extraction (no web framework)."""
from __future__ import annotations

import io
import os
import tempfile
import zipfile

ALLOWED_SUFFIXES = {'.txt', '.md', '.csv', '.docx', '.pdf'}
MIN_TEXT_CHARS = 50


class DocumentError(Exception):
    """A problem with an uploaded document. `status` is an HTTP-style code the web layer can reuse."""

    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


def validate_upload(data: bytes, suffix: str) -> None:
    if suffix not in ALLOWED_SUFFIXES:
        raise DocumentError(415, 'Supported files: PDF, DOCX, TXT, MD, CSV')
    if not data:
        raise DocumentError(400, 'The uploaded file is empty.')
    if suffix == '.pdf' and not data.startswith(b'%PDF'):
        raise DocumentError(415, 'File content does not match a PDF.')
    if suffix == '.docx':
        if not data.startswith(b'PK'):
            raise DocumentError(415, 'File content does not match a DOCX.')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                infos = z.infolist()
                if sum(i.file_size for i in infos) > 100_000_000 or len(infos) > 2000:
                    raise DocumentError(413, 'The DOCX expands to an unsafe size.')
        except zipfile.BadZipFile:
            raise DocumentError(415, 'File content does not match a DOCX.')


def extract_text(data: bytes, suffix: str, max_pages: int = 300) -> str:
    """Return the plain text of an upload. Raises DocumentError if it cannot be read or has no text."""
    validate_upload(data, suffix)
    try:
        if suffix in {'.txt', '.md', '.csv'}:
            text = data.decode('utf-8', errors='ignore')
        else:
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as t:
                t.write(data)
                path = t.name
            try:
                if suffix == '.docx':
                    from docx import Document as DocxDocument
                    d = DocxDocument(path)
                    text = '\n'.join(p.text for p in d.paragraphs)
                    for table in d.tables:
                        for row in table.rows:
                            text += '\n' + ' | '.join(c.text for c in row.cells)
                else:
                    from pypdf import PdfReader
                    text = '\n'.join((p.extract_text() or '') for p in PdfReader(path).pages[:max_pages])
            finally:
                os.unlink(path)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError(422, 'The document could not be read. It may be corrupt or password-protected.')
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise DocumentError(422, 'No readable text was found in this file. Scanned PDFs need OCR - please upload a '
                                 'text-based PDF or a DOCX.')
    return text
